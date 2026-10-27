"""Build the demo workbook: python src/build/build.py

Sheets (Arabic, right-to-left, like the production system):
  الإعدادات   settings, lists, chart of accounts, payroll rules
  اليومية      cashier's daily sales (5 channels) + drawer expenses + reconciliation
  الرواتب      payroll driven by settings and by advances taken from the drawer
  الدفتر       general ledger: manual invoices + automatic monthly postings
  قائمة الدخل  income statement by month from the ledger
  الضريبة      KSA VAT: output, input, due, paid, outstanding
  التحقق       cross-sheet validation checks
Every figure is a live formula; only the yellow cells are inputs.
"""
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

import data as D
from style import (ACCENT, CENTER, DATE, INT, NUM, PCT, SOFT, cell, font, header, setup,
                   title)

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "demo" / "finance-system-demo.xlsx"

S_SET, S_DAY, S_PAY, S_LED, S_PL, S_TAX, S_CHK = (
    "الإعدادات", "اليومية", "الرواتب", "الدفتر", "قائمة الدخل", "الضريبة", "التحقق")


def q(sheet):  # quoted sheet reference
    return f"'{sheet}'!"


SET, DAY, PAY, LED, PL = q(S_SET), q(S_DAY), q(S_PAY), q(S_LED), q(S_PL)
VAT_RATE = f"{SET}$N$4"


# ---------------------------------------------------------------- settings
def build_settings(wb):
    ws = wb.active
    ws.title = S_SET
    setup(ws, {"A": 16, "B": 3, "C": 14, "D": 12, "E": 12, "F": 3, "G": 9, "H": 30,
               "I": 18, "J": 10, "K": 14, "L": 3, "M": 24, "N": 12})
    title(ws, "A1:N1", "الإعدادات — القوائم ودليل الحسابات وقواعد الرواتب",
          "الخانات الصفراء فقط هي اللي بتتعدل — كل الشيتات التانية بتقرا منها")
    header(ws, 3, 1, ["أنواع مصروفات الدرج"])
    for i, t in enumerate(D.EXPENSE_TYPES):
        cell(ws, 4 + i, 1, t, inp=True)
    header(ws, 3, 3, ["الموظف", "الراتب الأساسي", "بدل المواصلات"])
    for i, (n, b, t) in enumerate(D.EMPLOYEES):
        cell(ws, 4 + i, 3, n, inp=True)
        cell(ws, 4 + i, 4, b, NUM, inp=True)
        cell(ws, 4 + i, 5, t, NUM, inp=True)
    header(ws, 3, 7, ["الكود", "الحساب", "التصنيف العام", "الطبيعة", "القائمة"])
    for i, row in enumerate(D.ACCOUNTS):
        for j, v in enumerate(row):
            cell(ws, 4 + i, 7 + j, v, inp=True)
    header(ws, 3, 13, ["القاعدة", "القيمة"])
    rules = [("نسبة ضريبة القيمة المضافة", 0.15, PCT), ("حد بداية العمولة (ريال مبيعات)", 3000, NUM),
             ("نسبة العمولة فوق الحد", 0.30, PCT), ("تارجت مكافأة المبيعات", 10000, NUM),
             ("قيمة مكافأة التارجت", 300, NUM), ("سعر ساعة الأوفر تايم", 20, NUM),
             ("أيام الشهر لحساب اليومية", 30, INT)]
    for i, (lab, v, f) in enumerate(rules):
        cell(ws, 4 + i, 13, lab, bold=True)
        cell(ws, 4 + i, 14, v, f, inp=True)
    ws.freeze_panes = "A4"


# ------------------------------------------------------------------- daily
DAY_FIRST = 5


def build_daily(wb, rows):
    ws = wb.create_sheet(S_DAY)
    widths = {"A": 11, "B": 9, "C": 12}
    for c in "DEFGHIJKLMNO":
        widths[c] = 9
    widths.update({"P": 11, "Q": 10, "R": 9, "S": 9, "T": 8, "U": 8, "V": 11, "W": 10,
                   "X": 9, "Y": 8, "Z": 8, "AA": 9, "AB": 11})
    setup(ws, widths)
    title(ws, "A1:AB1", "اليومية — مبيعات ومصروفات الدرج ومطابقة الكاش",
          "الكاشير بيدخل الخانات الصفراء بس — الإجماليات والمطابقة والصافي محسوبين")
    bands = [("A3:C3", "اليوم"), ("D3:P3", "مصروفات من الدرج"), ("Q3:V3", "المبيعات حسب طريقة الدفع (من النظام)"),
             ("W3:AA3", "مطابقة الفعلي"), ("AB3:AB3", "")]
    for rng, lab in bands:
        ws.merge_cells(rng)
        c = ws[rng.split(":")[0]]
        c.value = lab
        c.font = font(True, ACCENT)
        c.fill = PatternFill("solid", fgColor=SOFT)
        c.alignment = CENTER
    labels = ["التاريخ", "اليوم", "الشهر"]
    for k in range(1, 5):
        labels += [f"نوع {k}", f"الموظف {k}", f"القيمة {k}"]
    labels += ["إجمالي المصروفات", *D.CHANNELS, "إجمالي المبيعات", "الشبكة الفعلي",
               "الكاش الفعلي", "فرق الشبكة", "فرق الكاش", "الحالة", "صافي اليوم"]
    header(ws, 4, 1, labels)

    dv_type = DataValidation(type="list", formula1=f"={SET}$A$4:$A$7", allow_blank=True)
    dv_emp = DataValidation(type="list", formula1=f"={SET}$C$4:$C$9", allow_blank=True)
    ws.add_data_validation(dv_type)
    ws.add_data_validation(dv_emp)

    for i, r in enumerate(rows):
        n = DAY_FIRST + i
        cell(ws, n, 1, r["date"], DATE, inp=True)
        cell(ws, n, 2, r["day"])
        cell(ws, n, 3, r["month"])
        for k in range(4):
            col = 4 + k * 3
            if k < len(r["exps"]):
                t, e, a = r["exps"][k]
            else:
                t, e, a = None, None, None
            cell(ws, n, col, t, inp=True)
            cell(ws, n, col + 1, e or None, inp=True)
            cell(ws, n, col + 2, a, NUM, inp=True)
            dv_type.add(ws.cell(n, col))
            dv_emp.add(ws.cell(n, col + 1))
        cell(ws, n, 16, f"=F{n}+I{n}+L{n}+O{n}", NUM, bold=True)
        for j, ch in enumerate(D.CHANNELS):
            cell(ws, n, 17 + j, r["sales"][ch] or None, NUM, inp=True)
        cell(ws, n, 22, f"=SUM(Q{n}:U{n})", NUM, bold=True)
        cell(ws, n, 23, r["actual_network"], NUM, inp=True)
        cell(ws, n, 24, r["actual_cash"], NUM, inp=True)
        cell(ws, n, 25, f"=W{n}-Q{n}", NUM)
        cell(ws, n, 26, f"=X{n}-(R{n}-P{n})", NUM)  # drawer cash net of expenses paid from it
        cell(ws, n, 27, f'=IF(AND(Y{n}=0,Z{n}=0),"مطابق",IF(Y{n}+Z{n}>0,"زيادة","عجز"))')
        cell(ws, n, 28, f"=V{n}-P{n}", NUM, bold=True)
    last = DAY_FIRST + len(rows) - 1
    tot = last + 1
    cell(ws, tot, 1, "الإجمالي", bold=True, fill=SOFT)
    for col in [6, 9, 12, 15, 16, 17, 18, 19, 20, 21, 22, 23, 24, 25, 26, 28]:
        L = ws.cell(4, col).column_letter
        cell(ws, tot, col, f"=SUM({L}{DAY_FIRST}:{L}{last})", NUM, bold=True, fill=SOFT)
    ws.freeze_panes = "D5"
    return last


# ----------------------------------------------------------------- payroll
PAY_FIRST = 5


def build_payroll(wb, pay, day_last):
    ws = wb.create_sheet(S_PAY)
    setup(ws, {c: 11 for c in "ABCDEFGHIJKLMNOPQ"} | {"B": 13})
    title(ws, "A1:Q1", "الرواتب الشهرية",
          "العمولة والمكافأة والأوفر تايم من قواعد الإعدادات — السلف بتتسحب أوتوماتيك من اليومية")
    header(ws, 4, 1, ["السنة", "الشهر", "الموظف", "الأساسي", "المواصلات", "مبيعات الموظف",
                      "العمولة", "مكافأة التارجت", "ساعات أوفر تايم", "بدل الأوفر تايم",
                      "تسوية يدوية", "إجمالي المستحقات", "أيام غياب", "خصم الغياب",
                      "سلف من الدرج", "إجمالي الخصومات", "صافي الراتب"])
    rng = lambda col: f"{DAY}${col}${DAY_FIRST}:${col}${day_last}"
    for i, p in enumerate(pay):
        n = PAY_FIRST + i
        cell(ws, n, 1, p["year"])
        cell(ws, n, 2, p["month"])
        cell(ws, n, 3, p["name"], bold=True)
        cell(ws, n, 4, f"=INDEX({SET}$D$4:$D$9,MATCH(C{n},{SET}$C$4:$C$9,0))", NUM)
        cell(ws, n, 5, f"=INDEX({SET}$E$4:$E$9,MATCH(C{n},{SET}$C$4:$C$9,0))", NUM)
        cell(ws, n, 6, p["sales"], NUM, inp=True)
        cell(ws, n, 7, f"=MAX(0,F{n}-{SET}$N$5)*{SET}$N$6", NUM)
        cell(ws, n, 8, f"=IF(F{n}>={SET}$N$7,{SET}$N$8,0)", NUM)
        cell(ws, n, 9, p["ot_hours"] or None, INT, inp=True)
        cell(ws, n, 10, f"=I{n}*{SET}$N$9", NUM)
        cell(ws, n, 11, p["manual"] or None, NUM, inp=True)
        cell(ws, n, 12, f"=D{n}+E{n}+G{n}+H{n}+J{n}+K{n}", NUM, bold=True)
        cell(ws, n, 13, p["absence"] or None, INT, inp=True)
        cell(ws, n, 14, f"=M{n}*D{n}/{SET}$N$10", NUM)
        adv = "+".join(
            f'SUMIFS({rng(a)},{rng(t)},"سلفة",{rng(e)},$C{n},{DAY}$C${DAY_FIRST}:$C${day_last},$B{n})'
            for t, e, a in (("D", "E", "F"), ("G", "H", "I"), ("J", "K", "L"), ("M", "N", "O")))
        cell(ws, n, 15, "=" + adv, NUM)
        cell(ws, n, 16, f"=N{n}+O{n}", NUM)
        cell(ws, n, 17, f"=L{n}-P{n}", NUM, bold=True, fill=SOFT)
    last = PAY_FIRST + len(pay) - 1
    tot = last + 1
    cell(ws, tot, 1, "الإجمالي", bold=True, fill=SOFT)
    for col in range(4, 18):
        L = ws.cell(4, col).column_letter
        cell(ws, tot, col, f"=SUM({L}{PAY_FIRST}:{L}{last})", NUM, bold=True, fill=SOFT)
    ws.freeze_panes = "D5"
    return last


# ------------------------------------------------------------------ ledger
LED_FIRST = 5


def build_ledger(wb, manual, day_last, pay_last):
    ws = wb.create_sheet(S_LED)
    setup(ws, {"A": 11, "B": 34, "C": 26, "D": 22, "E": 17, "F": 11, "G": 11, "H": 9, "I": 7,
               "J": 9, "K": 18, "L": 7, "M": 15, "N": 8, "O": 12, "P": 6, "Q": 9, "R": 10,
               "S": 10, "T": 12})
    title(ws, "A1:T1", "الدفتر — كل الحركات بدليل حسابات مكوّد",
          "الصفوف الملوّنة قيود آلية من اليومية والرواتب — الكود والتصنيف والضريبة محسوبين لكل سطر")
    header(ws, 4, 1, ["التاريخ", "البيان", "الحساب", "المورد / الجهة", "الرقم الضريبي", "وارد",
                      "منصرف", "الحساب النقدي", "خاضع؟", "ربع الإقرار", "المصدر",
                      "الكود", "التصنيف العام", "الطبيعة", "القائمة", "الشهر", "الربع",
                      "ضريبة المخرجات", "ضريبة المدخلات", "الصافي قبل الضريبة"])

    dr = lambda col: f"{DAY}${col}${DAY_FIRST}:${col}${day_last}"
    pr = lambda col: f"{PAY}${col}${PAY_FIRST}:${col}${pay_last}"
    entries = [dict(d=r[0], desc=r[1], acc=r[2], party=r[3], vat=r[4], inn=None, out=r[5],
                    cash=r[6], vatable=r[7], retq=r[8], src=r[9] or "إدخال يدوي", auto=False)
               for r in manual]
    for y, m, label in D.MONTHS:
        end = date(y, m, 30 if m == 9 else 31)
        entries += [
            dict(d=end, desc=f"إيرادات {label} — من اليومية", acc="مبيعات الخدمات", party="",
                 vat="", inn=f'=SUMIFS({dr("V")},{dr("C")},"{label}")', out=None, cash="البنك",
                 vatable=1, retq="", src="قيد آلي — اليومية", auto=True),
            dict(d=end, desc=f"نثريات الدرج {label} (بدون السلف)", acc="مصروفات نثرية من الدرج",
                 party="", vat="", inn=None,
                 out=f'=SUMIFS({dr("P")},{dr("C")},"{label}")-SUMIFS({pr("O")},{pr("B")},"{label}")',
                 cash="الكاش", vatable=0, retq="", src="قيد آلي — اليومية", auto=True),
            dict(d=end, desc=f"رواتب وبدلات {label}", acc="الرواتب والبدلات", party="", vat="",
                 inn=None,
                 out=f'=SUMIFS({pr("D")},{pr("B")},"{label}")+SUMIFS({pr("E")},{pr("B")},"{label}")'
                     f'-SUMIFS({pr("N")},{pr("B")},"{label}")',
                 cash="البنك", vatable=0, retq="", src="قيد آلي — الرواتب", auto=True),
            dict(d=end, desc=f"عمولات وحوافز {label}", acc="العمولات والحوافز", party="", vat="",
                 inn=None,
                 out="=" + "+".join(f'SUMIFS({pr(c)},{pr("B")},"{label}")' for c in "GHJK"),
                 cash="البنك", vatable=0, retq="", src="قيد آلي — الرواتب", auto=True),
        ]
    entries.sort(key=lambda e: (e["d"], e["auto"]))

    for i, e in enumerate(entries):
        n = LED_FIRST + i
        fill = SOFT if e["auto"] else None
        inp = not e["auto"]
        cell(ws, n, 1, e["d"], DATE, inp=inp, fill=fill)
        cell(ws, n, 2, e["desc"], inp=inp, fill=fill)
        cell(ws, n, 3, e["acc"], inp=inp, fill=fill)
        cell(ws, n, 4, e["party"] or None, inp=inp, fill=fill)
        cell(ws, n, 5, e["vat"] or None, inp=inp, fill=fill)
        cell(ws, n, 6, e["inn"], NUM, inp=inp, fill=fill)
        cell(ws, n, 7, e["out"], NUM, inp=inp, fill=fill)
        cell(ws, n, 8, e["cash"], inp=inp, fill=fill)
        cell(ws, n, 9, e["vatable"], INT, inp=inp, fill=fill)
        cell(ws, n, 10, e["retq"] or None, inp=inp, fill=fill)
        cell(ws, n, 11, e["src"], fill=fill)
        look = lambda col: (f"=IFERROR(INDEX({SET}${col}$4:${col}$15,"
                            f"MATCH($C{n},{SET}$H$4:$H$15,0)),\"غير معرف\")")
        cell(ws, n, 12, look("G"))
        cell(ws, n, 13, look("I"))
        cell(ws, n, 14, look("J"))
        cell(ws, n, 15, look("K"))
        cell(ws, n, 16, f"=MONTH(A{n})")
        cell(ws, n, 17, f'=YEAR(A{n})&"-Q"&ROUNDUP(MONTH(A{n})/3,0)')
        cell(ws, n, 18, f"=IF(AND(I{n}=1,F{n}>0),F{n}*{VAT_RATE}/(1+{VAT_RATE}),0)", NUM)
        cell(ws, n, 19, f"=IF(AND(I{n}=1,G{n}>0),G{n}*{VAT_RATE}/(1+{VAT_RATE}),0)", NUM)
        cell(ws, n, 20, f"=F{n}+G{n}-R{n}-S{n}", NUM, bold=True)
    ws.freeze_panes = "C5"
    ws.auto_filter.ref = f"A4:T{LED_FIRST + len(entries) - 1}"
    return LED_FIRST + len(entries) - 1


# ------------------------------------------------------------- statements
def L_(col, last):
    return f"{LED}${col}${LED_FIRST}:${col}${last}"


def build_pl(wb, led_last):
    ws = wb.create_sheet(S_PL)
    setup(ws, {"A": 34, "B": 9, "C": 14, "D": 14, "E": 14, "F": 15})
    title(ws, "A1:F1", "قائمة الدخل — الربع الثالث 2026", "صافي قبل الضريبة — بتتحدث أوتوماتيك من الدفتر")
    for j, m in enumerate([7, 8, 9]):
        c = ws.cell(3, 3 + j, m)
        c.font = font(False, "9AA5A2", 8)
        c.alignment = CENTER
    header(ws, 4, 1, ["البند", "الكود", "يوليو", "أغسطس", "سبتمبر", "الإجمالي"])
    T, Lc, Pm = L_("T", led_last), L_("L", led_last), L_("P", led_last)

    def line(r, label, code, bold=False, fill=None):
        cell(ws, r, 1, label, bold=bold, fill=fill)
        cell(ws, r, 2, code, fill=fill)
        for j in range(3):
            col = ws.cell(3, 3 + j).column_letter
            cell(ws, r, 3 + j, f"=SUMIFS({T},{Lc},$B{r},{Pm},{col}$3)", NUM, bold, fill=fill)
        cell(ws, r, 6, f"=SUM(C{r}:E{r})", NUM, True, fill=fill)

    r = 5
    line(r, "مبيعات الخدمات", 4100)
    rev_row = r + 1
    cell(ws, rev_row, 1, "إجمالي الإيرادات", bold=True, fill=SOFT)
    cell(ws, rev_row, 2, None, fill=SOFT)
    for col in "CDEF":
        cell(ws, rev_row, " CDEF".index(col) + 2, f"={col}{r}", NUM, True, fill=SOFT)
    r = rev_row + 2
    exp_start = r
    for code, acc, _, nature, stmt in D.ACCOUNTS:
        if nature == "مصروف":
            line(r, acc, code)
            r += 1
    exp_end = r - 1
    cell(ws, r, 1, "إجمالي المصروفات", bold=True, fill=SOFT)
    cell(ws, r, 2, None, fill=SOFT)
    for k, col in enumerate("CDEF"):
        cell(ws, r, 3 + k, f"=SUM({col}{exp_start}:{col}{exp_end})", NUM, True, fill=SOFT)
    exp_tot = r
    r += 2
    cell(ws, r, 1, "صافي الربح", bold=True, fill="D7EEE9")
    cell(ws, r, 2, None, fill="D7EEE9")
    for k, col in enumerate("CDEF"):
        cell(ws, r, 3 + k, f"={col}{rev_row}-{col}{exp_tot}", NUM, True, fill="D7EEE9")
    net = r
    r += 1
    cell(ws, r, 1, "هامش صافي الربح", bold=True)
    cell(ws, r, 2, None)
    for k, col in enumerate("CDEF"):
        cell(ws, r, 3 + k, f"=IF({col}{rev_row}=0,0,{col}{net}/{col}{rev_row})", PCT, True)
    ws.freeze_panes = "C5"
    return net


def build_tax(wb, led_last):
    ws = wb.create_sheet(S_TAX)
    setup(ws, {"A": 22, "B": 16, "C": 15, "D": 16, "E": 15, "F": 15})
    title(ws, "A1:F1", "ضريبة القيمة المضافة — الربع الثالث 2026",
          "15% محسوبة على كل سطر في الدفتر — والمتبقي بيظهر كالتزام لحد ما يتسدد")
    header(ws, 4, 1, ["الفترة", "المبيعات الخاضعة", "ضريبة المخرجات", "المشتريات الخاضعة",
                      "ضريبة المدخلات", "صافي المستحق"])
    F_, G_, I_, P_ = (L_(c, led_last) for c in "FGIP")
    R_, S_ = L_("R", led_last), L_("S", led_last)
    for j, (y, m, label) in enumerate(D.MONTHS):
        n = 5 + j
        cell(ws, n, 1, label, bold=True)
        cell(ws, n, 2, f"=SUMIFS({F_},{I_},1,{P_},{m})", NUM)
        cell(ws, n, 3, f"=SUMIFS({R_},{P_},{m})", NUM)
        cell(ws, n, 4, f"=SUMIFS({G_},{I_},1,{P_},{m})", NUM)
        cell(ws, n, 5, f"=SUMIFS({S_},{P_},{m})", NUM)
        cell(ws, n, 6, f"=C{n}-E{n}", NUM, True)
    cell(ws, 8, 1, "إجمالي الربع", bold=True, fill=SOFT)
    for k, col in enumerate("BCDEF"):
        cell(ws, 8, 2 + k, f"=SUM({col}5:{col}7)", NUM, True, fill=SOFT)

    header(ws, 10, 1, ["الإقرار", "المستحق", "المسدد", "المتبقي (التزام)", "الحالة"])
    cell(ws, 11, 1, "2026-Q3", bold=True)
    cell(ws, 11, 2, "=F8", NUM, True)
    cell(ws, 11, 3, f'=SUMIFS({G_},{L_("L", led_last)},2300,{L_("J", led_last)},A11)', NUM)
    cell(ws, 11, 4, "=B11-C11", NUM, True, fill="FDECEA")
    cell(ws, 11, 5, '=IF(D11<=0,"مسدد بالكامل","سداد جزئي — التزام قائم")')


def build_checks(wb, day_last, pay_last, led_last, pl_net):
    ws = wb.create_sheet(S_CHK)
    setup(ws, {"A": 46, "B": 15, "C": 15, "D": 12, "E": 16})
    title(ws, "A1:E1", "التحقق — فحوصات قبل اعتماد أي تقرير",
          "✓ = سليم   •   ⚠ = يحتاج مراجعة")
    header(ws, 4, 1, ["الفحص", "القيمة الأولى", "القيمة التانية", "الفرق", "النتيجة"])
    Lr = lambda c: L_(c, led_last)
    day = lambda c: f"{DAY}{c}{DAY_FIRST}:{c}{day_last}"
    pay = lambda c: f"{PAY}{c}{PAY_FIRST}:{c}{pay_last}"
    checks = [
        ("مبيعات اليومية = إيرادات الدفتر", f"=SUM({day('V')})", f"=SUMIFS({Lr('F')},{Lr('L')},4100)"),
        ("مستحقات الرواتب − خصم الغياب = رواتب وحوافز الدفتر",
         f"=SUM({pay('L')})-SUM({pay('N')})",
         f"=SUMIFS({Lr('G')},{Lr('L')},5100)+SUMIFS({Lr('G')},{Lr('L')},5110)"),
        ("مصروفات الدرج = السلف + النثريات في الدفتر", f"=SUM({day('P')})",
         f"=SUM({pay('O')})+SUMIFS({Lr('G')},{Lr('L')},5900)"),
        ("حسابات غير معرفة في دليل الحسابات", f'=COUNTIF({Lr("L")},"غير معرف")', "=0"),
        ("صافي الربح في القائمة = صافي الدفتر", f"={PL}F{pl_net}",
         f'=SUMIFS({Lr("T")},{Lr("N")},"إيراد")-SUMIFS({Lr("T")},{Lr("N")},"مصروف")'),
        ("الضريبة المستحقة = المخرجات − المدخلات", f"={q(S_TAX)}F8",
         f"=SUM({Lr('R')})-SUM({Lr('S')})"),
        ("أيام فيها فرق كاش أو شبكة (للمراجعة)", f'=COUNTIF({day("AA")},"<>مطابق")', "=0"),
    ]
    for i, (lab, a, b) in enumerate(checks):
        n = 5 + i
        cell(ws, n, 1, lab, bold=True)
        cell(ws, n, 2, a, NUM)
        cell(ws, n, 3, b, NUM)
        cell(ws, n, 4, f"=B{n}-C{n}", NUM)
        cell(ws, n, 5, f'=IF(ABS(D{n})<0.01,"✓ سليم","⚠ راجع")', bold=True)


def main():
    daily, pay, manual = D.daily_rows(), D.payroll_inputs(), D.manual_ledger()
    D.export_csv(ROOT / "sample-data", daily, pay, manual)
    wb = Workbook()
    build_settings(wb)
    day_last = build_daily(wb, daily)
    pay_last = build_payroll(wb, pay, day_last)
    led_last = build_ledger(wb, manual, day_last, pay_last)
    pl_net = build_pl(wb, led_last)
    build_tax(wb, led_last)
    build_checks(wb, day_last, pay_last, led_last, pl_net)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"saved {OUT}")


if __name__ == "__main__":
    main()
