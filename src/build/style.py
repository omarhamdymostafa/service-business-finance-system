"""Shared visual style for every sheet in the demo workbook."""
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

FONT = "Arial"
ACCENT = "1F6F63"      # deep teal
ACCENT_2 = "3FA796"    # mid teal
SOFT = "E6F4F1"        # light teal (section bands)
INPUT = "FFF8E1"       # pale yellow = cells the user types in
OK = "2E7D32"
WARN = "C62828"

NUM = '#,##0.00;(#,##0.00);"-"'
INT = '#,##0;(#,##0);"-"'
PCT = '0.0%'
DATE = 'yyyy/mm/dd'

thin = Side(style="thin", color="D5DEDC")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def font(bold=False, color="222222", size=10):
    return Font(name=FONT, bold=bold, color=color, size=size)


def setup(ws, widths):
    """Right-to-left sheet, default widths, gridlines off."""
    ws.sheet_view.rightToLeft = True
    ws.sheet_view.showGridLines = False
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def title(ws, rng, text, sub=None):
    first = rng.split(":")[0]
    ws.merge_cells(rng)
    c = ws[first]
    c.value = text
    c.font = font(True, "FFFFFF", 14)
    c.fill = PatternFill("solid", fgColor=ACCENT)
    c.alignment = CENTER
    ws.row_dimensions[c.row].height = 30
    if sub:
        r = c.row + 1
        start, end = rng.split(":")
        sub_rng = f"{start[0:len(start)-len(str(c.row))]}{r}:{end[0:len(end)-len(str(c.row))]}{r}"
        ws.merge_cells(sub_rng)
        s = ws[sub_rng.split(":")[0]]
        s.value = sub
        s.font = font(False, "4A5653", 9)
        s.alignment = CENTER


def header(ws, row, col_start, labels, fill=ACCENT_2):
    for i, lab in enumerate(labels):
        c = ws.cell(row=row, column=col_start + i, value=lab)
        c.font = font(True, "FFFFFF")
        c.fill = PatternFill("solid", fgColor=fill)
        c.alignment = CENTER
        c.border = BORDER
    ws.row_dimensions[row].height = 32


def cell(ws, row, col, value, fmt=None, bold=False, inp=False, fill=None, color="222222"):
    c = ws.cell(row=row, column=col, value=value)
    c.font = font(bold, color)
    c.border = BORDER
    c.alignment = Alignment(horizontal="center", vertical="center")
    if fmt:
        c.number_format = fmt
    if inp:
        c.fill = PatternFill("solid", fgColor=INPUT)
    elif fill:
        c.fill = PatternFill("solid", fgColor=fill)
    return c
